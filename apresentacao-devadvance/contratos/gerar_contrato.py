"""Gera contrato preenchido (docx + pdf) a partir do template.
python gerar_contrato.py --nome "..." --cpf "..." --endereco "..." --telefone "..." --email "..." --cidade "..." --valor 199 --plano mensal [--taxa 250 --parcelas-taxa 3]
"""
import argparse, copy, datetime, pathlib, subprocess
from docx import Document
from dinheiro import brl, extenso, split_cents

HERE = pathlib.Path(__file__).resolve().parent
TEMPLATE = HERE / "Contrato_Programa_ADVANCE_Turma_Fundadora.docx"
SOFFICE = "/mnt/c/Program Files/LibreOffice/program/soffice.exe"
MESES = ["janeiro","fevereiro","março","abril","maio","junho","julho","agosto","setembro","outubro","novembro","dezembro"]

ap = argparse.ArgumentParser()
for k in ["nome","cpf","endereco","telefone","email","cidade"]:
    ap.add_argument(f"--{k}", required=True)
ap.add_argument("--valor", type=float, required=True, help="mensal: valor do mês; anual: valor total")
ap.add_argument("--plano", required=True, choices=["mensal", "anual"])
ap.add_argument("--taxa", type=float, default=0, help="taxa de adesão em reais")
ap.add_argument("--parcelas-taxa", type=int, default=1, choices=[1, 2, 3])
a = ap.parse_args()
assert a.plano == "mensal" or a.parcelas_taxa == 1, "anual: taxa vai inteira na cobrança única"
rs = lambda v: f"R$ {brl(v)} ({extenso(v)})"
PARC = {1: "1 (uma)", 2: "2 (duas)", 3: "3 (três)"}

hoje = datetime.date.today()
subs = {
    "[NOME COMPLETO]": a.nome, "[CPF]": a.cpf, "[ENDEREÇO COMPLETO]": a.endereco,
    "[TELEFONE]": a.telefone, "[E-MAIL]": a.email,
    "[CIDADE]": a.cidade, "[DATA]": f"{hoje.day} de {MESES[hoje.month-1]} de {hoje.year}",
}

d = Document(TEMPLATE)
# Cláusula 05.1: só a modalidade contratada, marcada, com o valor do argumento
ps = d.paragraphs
i = next(i for i, p in enumerate(ps) if p.text.startswith("05.1 –"))
ps[i].runs[0].text = "05.1 – O CONTRATANTE optou pela seguinte modalidade de pagamento:"
mensal, anual = ps[i + 1], ps[i + 2]
assert mensal.text.startswith("[  ] Mensal:") and anual.text.startswith("[  ] Anual"), "template mudou"
if a.plano == "mensal":
    esc, fora = mensal, anual
    esc.runs[0].text = "[X] Mensal:"
    esc.runs[1].text = f" {rs(a.valor)} por mês, por cobrança recorrente em cartão de crédito ou boleto bancário, com vencimento na mesma data de cada mês;"
else:
    esc, fora = anual, mensal
    esc.runs[0].text = "[X] Anual à vista:"
    esc.runs[1].text = f" {rs(a.valor)}, correspondente aos 12 (doze) meses do Programa;"
fora._p.getparent().remove(fora._p)
if a.taxa:
    n, parc = a.parcelas_taxa, split_cents(a.taxa, a.parcelas_taxa)
    if n == 1:
        como = "paga integralmente junto com a " + ("primeira mensalidade." if a.plano == "mensal" else "cobrança anual.")
    elif parc[0] == parc[-1]:
        como = f"paga em {PARC[n]} parcelas mensais de {rs(parc[0])}, somadas às {PARC[n].split()[0]} primeiras mensalidades."
    else:
        como = (f"paga em {PARC[n]} parcelas mensais, somadas às {PARC[n].split()[0]} primeiras mensalidades, "
                f"sendo a primeira de {rs(parc[0])} e as demais de {rs(parc[-1])}.")
    taxa = copy.deepcopy(esc._p)
    esc._p.addnext(taxa)
    taxa = next(p for p in d.paragraphs if p._p is taxa)
    taxa.runs[0].text = "Taxa de adesão:"
    taxa.runs[1].text = f" {rs(a.taxa)}, {como}"
# 09.4: meses iniciados do anual descontados a 1/10 do valor (12 pelo preço de 10)
ref = a.valor if a.plano == "mensal" else round(a.valor / 10, 2)
subs["R$ 197,00 (cento e noventa e sete reais)"] = rs(ref)

for p in d.paragraphs:
    for r in p.runs:
        for k, v in subs.items():
            if k in r.text:
                r.text = r.text.replace(k, v)
    if p.text.strip() == "CONTRATANTE":  # bloco de assinatura
        p.add_run(f"\n{a.nome}\nCPF: {a.cpf}")

out = HERE / f"Contrato_Programa_ADVANCE_{a.nome}.docx"
d.save(out)
win = lambda p: subprocess.check_output(["wslpath", "-w", str(p)]).decode().strip()
subprocess.run([SOFFICE, "--headless", "--convert-to", "pdf", "--outdir", win(HERE), win(out)], check=True)
pdf = out.with_suffix(".pdf")

# posições (%) pra campos de assinatura no Autentique: 1 acima de cada label do bloco de assinatura
import fitz, json
pos = {}
for i, pg in enumerate(fitz.open(pdf)):
    W, H = pg.rect.width, pg.rect.height
    for b in pg.get_text("blocks"):
        t = b[4].strip()
        for k in ("CONTRATANTE", "CONTRATADA"):
            if t.startswith(k) and not t.startswith(k + ":"):
                pos[k] = {"z": i + 1, "x": str(round(b[0] / W * 100, 1)), "y": str(round(b[1] / H * 100 - 7, 1))}
print(out)
print(pdf)
print("SIZE", pdf.stat().st_size)
print("POSICOES", json.dumps(pos))
