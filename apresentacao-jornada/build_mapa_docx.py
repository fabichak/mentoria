"""Gera mapa-da-jornada.docx: exercício do workshop, editável (sobe direto no Google Docs).
python build_mapa_docx.py"""
import pathlib
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = pathlib.Path(__file__).with_name("mapa-da-jornada.docx")
CYAN, MUTED, INK = RGBColor(0x08,0x91,0xb2), RGBColor(0x5b,0x6b,0x85), RGBColor(0x0b,0x12,0x20)

PASSOS = [
  ("1. ONDE ESTOU",        "cargo, salário, força, lacuna",
     ["Cargo atual", "Salário atual", "Minha maior força", "Minha maior lacuna"]),
  ("2. MINHA TRILHA",      "Dev / Liderança / Empreendedor / explorando",
     ["Trilha que escolho (ou estou explorando)", "Por quê"]),
  ("3. DESTINO 24 MESES",  "cargo + faixa salarial",
     ["Cargo em 24 meses", "Faixa salarial", "Onde (empresa, país, formato)"]),
  ("4. 2 ARMADILHAS",      "das 5, quais me pegam hoje: auto-estima · auto-sabotagem · falta de network · só técnico · falta de comunidade",
     ["Armadilha 1", "Como ela aparece no meu dia a dia", "Armadilha 2", "Como ela aparece no meu dia a dia"]),
  ("5. PRÓXIMOS 3 PASSOS", "1 esta semana · 1 este mês · 1 este trimestre",
     ["Esta semana", "Este mês", "Este trimestre"]),
]

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2); s.left_margin = s.right_margin = Cm(2.2)
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(11); st.font.color.rgb = INK

def para(text, size=11, color=INK, bold=False, mono=False, space_after=4, space_before=0):
    p = doc.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    if mono: r.font.name = 'Courier New'
    p.paragraph_format.space_after = Pt(space_after); p.paragraph_format.space_before = Pt(space_before)
    return p

def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_); tcPr.append(shd)

def campos(lista, linhas=3):
    # tabela 2 colunas: rótulo | espaço pra escrever (linhas vazias = altura)
    t = doc.add_table(rows=len(lista), cols=2); t.style = 'Table Grid'
    for row, label in zip(t.rows, lista):
        a, b = row.cells
        a.width, b.width = Cm(5.2), Cm(11.4)
        a.paragraphs[0].add_run(label).bold = True; shade(a, 'EEF3FA')
        b.paragraphs[0].text = ''
        for _ in range(linhas - 1): b.add_paragraph('')
    doc.add_paragraph()

# capa
para('MENTORIA · DEVADVANCE.CLUB', 9, CYAN, mono=True)
para('Seu Mapa da Jornada', 26, INK, True, space_after=2)
para('Exercício guiado do workshop A Jornada do Desenvolvedor.', 12, MUTED)
para('Preencha os 5 passos e guarde. Volte nele a cada 90 dias.', 12, CYAN, True, space_after=14)
para('Como usar: Arquivo → Fazer uma cópia (Google Docs) e escreva direto nas caixas. Não precisa ficar bonito, precisa ficar honesto.', 10, MUTED, space_after=6)

for k, d, lista in PASSOS:
    doc.add_page_break()
    para('MAPA DA JORNADA', 9, CYAN, mono=True)
    para(k, 18, INK, True, space_after=2)
    para(d, 11, MUTED, space_after=10)
    campos(lista, linhas=3 if len(lista) <= 3 else 2)

doc.add_page_break()
para('MAPA DA JORNADA · RESUMO', 9, CYAN, mono=True)
para('Meu mapa em uma página', 18, INK, True, space_after=2)
para('Copie aqui o essencial de cada passo. Esta é a página que você revisita.', 11, MUTED, space_after=10)
campos([f'{k}\n{d.split(":")[0]}' for k, d, _ in PASSOS], linhas=3)

doc.save(OUT); print('OK:', OUT.name)
